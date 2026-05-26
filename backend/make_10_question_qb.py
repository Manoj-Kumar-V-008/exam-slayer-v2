from docx import Document

def main():
    doc = Document()
    doc.add_heading("Database Management Systems - 10 Question Bank", level=1)
    
    # 1. Define DBMS (2 Marks, definition)
    doc.add_paragraph("1. Define DBMS and state its basic purpose. (2 Marks)")
    
    # 2. Three-schema architecture (10 Marks, theory)
    doc.add_paragraph("2. Explain the three-schema architecture with a diagram showing physical, logical, and view levels. Discuss how physical and logical data independence are achieved. [10 Marks]")
    
    # 3. Relational vs NoSQL (5 Marks, comparison)
    doc.add_paragraph("3) Differentiate between Relational DBMS (RDBMS) and NoSQL databases. (5m)")
    
    # 4. SQL queries (10 Marks, sql)
    doc.add_paragraph("Question 4 (10 marks): Consider the following database schema:")
    doc.add_paragraph("Student(student_id, student_name, major)")
    doc.add_paragraph("Course(course_id, course_name, credits)")
    doc.add_paragraph("Enrollment(student_id, course_id, grade)")
    doc.add_paragraph("Write SQL queries to:")
    doc.add_paragraph("a) Retrieve all student names.")
    doc.add_paragraph("b) Count the total number of students.")
    doc.add_paragraph("c) Find students enrolled in Course 'CS101' with grade 'A'.")
    
    # 5. Relational algebra (5 Marks, relational_algebra)
    doc.add_paragraph("5. Illustrate the relational algebra operations for select and project. Write a relational algebra expression to retrieve students majoring in 'Computer Science' from the Student relation. [5 Marks]")
    
    # 6. Candidate Key (2 Marks, theory)
    doc.add_paragraph("Q6. What is a Candidate Key? Explain with a short example. (2 Marks)")
    
    # 7. ER to Relational mapping (10 Marks, er_model)
    doc.add_paragraph("7. Given an ER model with Student and Course entities in a Many-to-Many relationship, explain how to convert this ER model into a relational schema. Identify primary keys and foreign keys for the resulting relations. (10 marks)")
    
    # 8. ACID Properties (5 Marks, theory)
    doc.add_paragraph("8. Explain ACID properties in DBMS with a brief example for each property. [5 Marks]")
    
    # 9. Referential Integrity (2 Marks, definition)
    doc.add_paragraph("Question 9: Define the term Referential Integrity Constraint. (2 Marks)")
    
    # 10. Normalization problem (5 Marks, problem_solving)
    doc.add_paragraph("10) Consider a relation R(A, B, C, D) with functional dependencies A -> B and B -> C. Walk through the normalization process to convert this relation into 3NF step-by-step. (5 Marks)")
    
    output_path = "test_10_question_bank.docx"
    doc.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == "__main__":
    main()
